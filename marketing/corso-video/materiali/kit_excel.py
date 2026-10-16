"""Kit Excel del corso di gestione del rischio (Riskly). Genera Riskly-kit-rischio.xlsx.
Celle gialle con testo blu = da modificare; tutto il resto sono formule."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

F = 'Arial'
VERDE = '15181B'  # inchiostro Riskly
f_tit = Font(name=F, size=16, bold=True, color='B97C00')
f_sot = Font(name=F, size=10, italic=True, color='666666')
f_int = Font(name=F, size=10, bold=True, color='FFFFFF')
f_lab = Font(name=F, size=10)
f_in = Font(name=F, size=10, color='0000FF', bold=True)
f_out = Font(name=F, size=10, bold=True)
f_big = Font(name=F, size=13, bold=True, color='B97C00')
fill_in = PatternFill('solid', fgColor='FFF4B0')
fill_int = PatternFill('solid', fgColor=VERDE)
fill_out = PatternFill('solid', fgColor='EEF0EC')
sottile = Side(style='thin', color='CCCCCC')
bordo = Border(left=sottile, right=sottile, top=sottile, bottom=sottile)
EUR = '#,##0.00 "€";-#,##0.00 "€";"-"'
EUR0 = '#,##0 "€";-#,##0 "€";"-"'
PCT = '0.0%;-0.0%;"-"'
PCT2 = '0.00%'
AVVISO = 'Materiale educativo del corso Riskly, non consulenza finanziaria. Il trading con leva può farti perdere tutto il capitale.'

wb = Workbook()

def foglio(nome, titolo, sotto, larg):
    ws = wb.create_sheet(nome)
    ws.sheet_view.showGridLines = False
    for col, w in larg.items():
        ws.column_dimensions[col].width = w
    ws['A1'] = titolo; ws['A1'].font = f_tit
    ws['A2'] = sotto; ws['A2'].font = f_sot
    return ws

def lab(ws, c, t, font=f_lab):
    ws[c] = t; ws[c].font = font

def inp(ws, c, v, fmt=None, nota=None):
    ws[c] = v; ws[c].font = f_in; ws[c].fill = fill_in; ws[c].border = bordo
    if fmt: ws[c].number_format = fmt
    if nota: ws[c].comment = Comment(nota, 'Riskly')

def out(ws, c, formula, fmt=None, big=False):
    ws[c] = formula; ws[c].font = f_big if big else f_out; ws[c].fill = fill_out; ws[c].border = bordo
    if fmt: ws[c].number_format = fmt

def intest(ws, riga, col0, nomi):
    for i, n in enumerate(nomi):
        c = ws.cell(row=riga, column=col0 + i, value=n)
        c.font = f_int; c.fill = fill_int; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = bordo

def avviso(ws, riga):
    ws.cell(row=riga, column=1, value=AVVISO).font = f_sot

# ---------- Leggimi ----------
ws = wb.active; ws.title = 'Leggimi'; ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 28; ws.column_dimensions['B'].width = 90
ws['A1'] = 'Riskly · Kit del corso di gestione del rischio'; ws['A1'].font = f_tit
ws['A2'] = 'Sei fogli, uno per ogni strumento usato nelle lezioni. Funziona con Excel, Google Fogli e LibreOffice.'; ws['A2'].font = f_sot
ws['A4'] = 'Come si usa'; ws['A4'].font = f_out
inp(ws, 'A5', 'Cella gialla, testo blu'); ws['B5'] = 'Da modificare con i tuoi numeri. Le celle con il triangolino rosso hanno una nota che spiega cosa scrivere.'
out(ws, 'A6', 'Cella verde'); ws['B6'] = 'Risultato: è una formula, non va toccata.'
righe = [
    ('Calcolatore lotto', 'Lezioni 1.2–1.5: dal rischio in % allo stop in pip, al lotto giusto.'),
    ('Piano di rischio', 'Lezioni 6.x e 10.x: limiti giornalieri, settimanali e mensili in euro e in numero di stop.'),
    ('Recupero drawdown', 'Lezione 3.x: quanto devi guadagnare per tornare in pari dopo una perdita.'),
    ('Serie di perdite', 'Moduli 3–4: cosa fa al conto una serie di stop di fila e quanto è probabile.'),
    ('Expectancy', 'Lezione 2.1: quanto vale in media un trade del tuo sistema, in R e in euro.'),
    ('Journal', 'Lezioni 8.x: registro delle operazioni con risultato in R e statistiche automatiche.'),
]
ws['A8'] = 'I fogli'; ws['A8'].font = f_out
for i, (n, d) in enumerate(righe):
    ws.cell(row=9 + i, column=1, value=n).font = Font(name=F, size=10, bold=True, color=VERDE)
    ws.cell(row=9 + i, column=2, value=d).font = f_lab
ws['A16'] = 'Ipotesi'; ws['A16'].font = f_out
ws['B16'] = 'Valore del pip: 10 € per lotto standard (EUR/USD con conto in euro, valore approssimato). Per altri strumenti scrivi il valore del tuo broker.'
ws['B17'] = 'Probabilità delle serie di perdite: ogni trade indipendente dagli altri, win rate costante. Nella realtà i risultati possono essere peggiori.'
for r in (16, 17): ws.cell(row=r, column=2).font = f_lab
ws['A19'] = AVVISO; ws['A19'].font = f_sot
for r in range(5, 7): ws.cell(row=r, column=2).font = f_lab

# ---------- Calcolatore lotto ----------
ws = foglio('Calcolatore lotto', 'Calcolatore lotto', 'Prima scegli lo stop dove l\'idea è sbagliata, poi il lotto si adatta al rischio.', {'A': 34, 'B': 18, 'C': 4, 'D': 60})
lab(ws, 'A4', 'Capitale del conto (€)'); inp(ws, 'B4', 10000, EUR0, 'Il saldo attuale del conto.')
lab(ws, 'A5', 'Rischio per trade (%)'); inp(ws, 'B5', 0.01, PCT2, 'Scrivi 1% per l\'1%. Riferimento comune: 0,5%–1%.')
lab(ws, 'A6', 'Stop loss (pip)'); inp(ws, 'B6', 25, '0.0', 'Distanza tra prezzo di ingresso e stop, in pip.')
lab(ws, 'A7', 'Valore di 1 pip per 1 lotto (€)'); inp(ws, 'B7', 10, EUR, 'EUR/USD con conto in euro: circa 10 € (varia col cambio). Oro e indici: chiedi al broker.')
lab(ws, 'A8', 'Lotto minimo del broker'); inp(ws, 'B8', 0.01, '0.00', 'Di solito 0,01 (micro lotto).')
lab(ws, 'A10', 'Rischio in euro', f_out); out(ws, 'B10', '=B4*B5', EUR)
lab(ws, 'A11', 'Lotto esatto'); out(ws, 'B11', '=IF(B6*B7=0,0,B10/(B6*B7))', '0.0000')
lab(ws, 'A12', 'Lotto da usare (arrotondato per difetto)', f_out); out(ws, 'B12', '=IF(B8=0,B11,FLOOR(B11+0.0000001,B8))', '0.00', big=True)
lab(ws, 'A13', 'Perdita reale se scatta lo stop'); out(ws, 'B13', '=B12*B6*B7', EUR)
lab(ws, 'A14', 'Rischio reale in % del conto'); out(ws, 'B14', '=IF(B4=0,0,B13/B4)', PCT2)
ws['D10'] = 'Formula: lotto = rischio in € ÷ (stop in pip × valore del pip)'
ws['D12'] = 'Si arrotonda sempre per difetto: meglio rischiare un po\' meno che un po\' di più.'
ws['D14'] = 'Se è 0 il lotto minimo è troppo grande per questo stop: allarga il conto o salta il trade.'
for c in ('D10', 'D12', 'D14'): ws[c].font = f_sot
ws['A17'] = 'Stessa operazione, stop diversi'; ws['A17'].font = f_out
intest(ws, 18, 1, ['Stop (pip)', 'Lotto'])
for i, s in enumerate([10, 15, 20, 25, 30, 40, 50, 75, 100]):
    r = 19 + i
    inp(ws, f'A{r}', s, '0')
    out(ws, f'B{r}', f'=IF($B$8=0,$B$10/(A{r}*$B$7),FLOOR($B$10/(A{r}*$B$7)+0.0000001,$B$8))', '0.00')
avviso(ws, 29)

# ---------- Piano di rischio ----------
ws = foglio('Piano di rischio', 'Il mio piano di rischio', 'Le regole si scrivono prima, a mercato chiuso. Qui diventano numeri.', {'A': 40, 'B': 16, 'C': 18, 'D': 22})
lab(ws, 'A4', 'Capitale del conto (€)'); inp(ws, 'B4', 10000, EUR0)
lab(ws, 'A5', 'Rischio per trade (%)'); inp(ws, 'B5', 0.01, PCT2)
lab(ws, 'A6', 'Perdita massima al giorno (%)'); inp(ws, 'B6', 0.03, PCT2, 'Raggiunta questa perdita smetti per oggi.')
lab(ws, 'A7', 'Perdita massima alla settimana (%)'); inp(ws, 'B7', 0.06, PCT2)
lab(ws, 'A8', 'Perdita massima al mese (%)'); inp(ws, 'B8', 0.10, PCT2, 'Superata: stop, revisione del sistema e rischio dimezzato.')
lab(ws, 'A9', 'Operazioni massime al giorno'); inp(ws, 'B9', 3, '0')
lab(ws, 'A10', 'Riduzione del rischio dopo il limite mensile'); inp(ws, 'B10', 0.5, '0%', '50% = rischio dimezzato.')
intest(ws, 12, 1, ['Limite', 'In euro', 'In numero di stop', 'Controllo'])
lim = [('Per trade', '=B4*B5', None), ('Al giorno', '=B4*B6', 'B6'), ('Alla settimana', '=B4*B7', 'B7'), ('Al mese', '=B4*B8', 'B8')]
for i, (n, f, p) in enumerate(lim):
    r = 13 + i
    ws.cell(row=r, column=1, value=n).font = f_lab
    out(ws, f'B{r}', f, EUR)
    out(ws, f'C{r}', '=1' if p is None else f'=IF($B$5=0,0,ROUNDDOWN({p}/$B$5,0))', '0')
ws['D13'] = '=IF(B5>0.02,"Rischio alto: oltre il 2%","OK")'
ws['D14'] = '=IF(C14<B9,"Ti fermi prima di finire i trade del giorno","OK")'
ws['D15'] = '=IF(B7<B6,"Il limite settimanale è sotto quello giornaliero","OK")'
ws['D16'] = '=IF(B8<B7,"Il limite mensile è sotto quello settimanale","OK")'
for r in range(13, 17):
    ws[f'D{r}'].font = f_out; ws[f'D{r}'].border = bordo
ws.conditional_formatting.add('D13:D16', CellIsRule(operator='notEqual', formula=['"OK"'], font=Font(name=F, bold=True, color='C0392B')))
lab(ws, 'A18', 'Rischio per trade dopo il limite mensile'); out(ws, 'B18', '=B5*B10', PCT2)
lab(ws, 'A19', 'In euro'); out(ws, 'B19', '=B4*B18', EUR)
ws['A21'] = 'Le mie regole (scrivile tu)'; ws['A21'].font = f_out
for i, t in enumerate(['Quando opero:', 'Su quali strumenti:', 'Quando NON opero (notizie, stanchezza, dopo 2 stop…):', 'Cosa faccio dopo una perdita:', 'Cosa faccio dopo una vincita:']):
    ws.cell(row=22 + i, column=1, value=t).font = f_lab
    ws.merge_cells(start_row=22 + i, start_column=2, end_row=22 + i, end_column=4)
    c = ws.cell(row=22 + i, column=2); c.fill = fill_in; c.font = f_in; c.border = bordo
ws['B22'] = 'es. 9:00–12:00, mai nei 15 minuti prima e dopo un dato forte'
avviso(ws, 28)

# ---------- Recupero drawdown ----------
ws = foglio('Recupero drawdown', 'Recupero del drawdown', 'Più perdi, più è difficile tornare in pari: la salita è sempre più ripida della discesa.', {'A': 22, 'B': 22, 'C': 26, 'D': 4, 'E': 34, 'F': 16})
intest(ws, 4, 1, ['Perdita (drawdown)', 'Guadagno per tornare in pari', 'Trade vincenti di fila (guadagno in F6)'])
for i, d in enumerate([0.05, 0.10, 0.15, 0.20, 0.25, 0.30, 0.40, 0.50, 0.60, 0.75, 0.90]):
    r = 5 + i
    inp(ws, f'A{r}', d, '0%')
    out(ws, f'B{r}', f'=A{r}/(1-A{r})', '0.0%')
    out(ws, f'C{r}', f'=ROUNDUP(LN(1/(1-A{r}))/LN(1+$F$6),0)', '0')
ws['E4'] = 'Il tuo caso'; ws['E4'].font = f_out
lab(ws, 'E5', 'Drawdown attuale (%)'); inp(ws, 'F5', 0.12, '0.0%')
lab(ws, 'E6', 'Guadagno per trade vincente (%)'); inp(ws, 'F6', 0.01, '0.0%', 'Usato anche nella colonna C. 1% = un trade da +1R rischiando l\'1%.')
lab(ws, 'E7', 'Guadagno per tornare in pari', f_out); out(ws, 'F7', '=F5/(1-F5)', '0.0%', big=True)
lab(ws, 'E8', 'Trade vincenti di fila necessari'); out(ws, 'F8', '=ROUNDUP(LN(1/(1-F5))/LN(1+F6),0)', '0')
ws['E10'] = 'Formula: recupero = perdita ÷ (1 − perdita)'; ws['E10'].font = f_sot
ws['E11'] = 'Trade vincenti calcolati con interesse composto, senza perdite in mezzo.'; ws['E11'].font = f_sot
avviso(ws, 18)

# ---------- Serie di perdite ----------
ws = foglio('Serie di perdite', 'Serie di perdite', 'Cosa fa al conto una serie di stop di fila, e quanto è probabile vederla.', {'A': 16, 'B': 18, 'C': 16, 'D': 22, 'E': 4, 'F': 40, 'G': 14})
lab(ws, 'F4', 'Capitale (€)'); inp(ws, 'G4', 10000, EUR0)
lab(ws, 'F5', 'Rischio per trade (%)'); inp(ws, 'G5', 0.01, PCT2)
lab(ws, 'F6', 'Win rate del sistema (%)'); inp(ws, 'G6', 0.45, '0%', 'Sistema d\'esempio del corso: 45%.')
lab(ws, 'F7', 'Numero di trade (anno)'); inp(ws, 'G7', 200, '0')
lab(ws, 'F9', 'Serie di perdite più lunga attesa', f_out)
out(ws, 'G9', '=IF(OR(G6<=0,G6>=1),0,ROUND(LN(G7*G6)/LN(1/(1-G6)),0))', '0', big=True)
ws['F10'] = 'Stima classica: ln(trade × win rate) ÷ ln(1 ÷ % perdite).'; ws['F10'].font = f_sot
ws['F11'] = 'Con 45% e 200 trade è normale vedere 8–9 stop di fila.'; ws['F11'].font = f_sot
intest(ws, 4, 1, ['Stop di fila', 'Capitale dopo', 'Drawdown', 'Probabilità di vederla almeno una volta'])
for n in range(1, 21):
    r = 4 + n
    ws.cell(row=r, column=1, value=n).font = f_lab
    ws.cell(row=r, column=1).alignment = Alignment(horizontal='center')
    out(ws, f'B{r}', f'=$G$4*(1-$G$5)^A{r}', EUR0)
    out(ws, f'C{r}', f'=1-(1-$G$5)^A{r}', '0.0%')
    # stima di Feller semplificata: P(almeno una serie di n perdite in N trade) ≈ 1-(1-p*q^n)^(N-n+1)
    out(ws, f'D{r}', f'=IF($G$7<A{r},0,MIN(1,1-(1-$G$6*(1-$G$6)^A{r})^($G$7-A{r}+1)))', '0.0%')
ws.conditional_formatting.add('D5:D24', CellIsRule(operator='greaterThanOrEqual', formula=['0.5'], font=Font(name=F, bold=True, color='C0392B')))
ws['A26'] = 'Probabilità: approssimazione con trade indipendenti (una serie inizia dopo una vincita). In rosso = più probabile che no.'; ws['A26'].font = f_sot
avviso(ws, 27)

# ---------- Expectancy ----------
ws = foglio('Expectancy', 'Expectancy del sistema', 'Quanto vale, in media, un trade del tuo sistema. Si misura in R.', {'A': 36, 'B': 18, 'C': 4, 'D': 60})
lab(ws, 'A4', 'Win rate (%)'); inp(ws, 'B4', 0.45, '0%')
lab(ws, 'A5', 'Guadagno medio quando vinci (R)'); inp(ws, 'B5', 1.8, '0.00')
lab(ws, 'A6', 'Perdita media quando perdi (R)'); inp(ws, 'B6', 1, '0.00', 'Di solito 1 se rispetti sempre lo stop.')
lab(ws, 'A7', 'Valore di 1R in euro'); inp(ws, 'B7', 100, EUR0, 'Il rischio per trade in euro (es. 1% di 10.000 €).')
lab(ws, 'A8', 'Trade previsti'); inp(ws, 'B8', 100, '0')
lab(ws, 'A10', 'Expectancy (R per trade)', f_out); out(ws, 'B10', '=B4*B5-(1-B4)*B6', '+0.00;-0.00;0.00', big=True)
lab(ws, 'A11', 'In euro per trade'); out(ws, 'B11', '=B10*B7', EUR)
lab(ws, 'A12', 'Risultato atteso sui trade previsti'); out(ws, 'B12', '=B10*B7*B8', EUR0)
lab(ws, 'A13', 'Win rate di pareggio'); out(ws, 'B13', '=IF(B5+B6=0,0,B6/(B5+B6))', '0.0%')
lab(ws, 'A14', 'Giudizio'); out(ws, 'B14', '=IF(B10>0,"Vantaggio","Nessun vantaggio")')
ws['D10'] = 'E = W × guadagno medio − (1 − W) × perdita media'
ws['D12'] = 'È una media: il risultato reale di 100 trade può essere molto diverso.'
ws['D13'] = 'Sotto questo win rate, con questi guadagni e perdite, il sistema perde.'
ws['D14'] = 'Se non c\'è vantaggio, nessuna gestione del rischio lo salva.'
for c in ('D10', 'D12', 'D13', 'D14'): ws[c].font = f_sot
avviso(ws, 17)

# ---------- Journal ----------
ws = foglio('Journal', 'Journal delle operazioni', 'Una riga per trade. La prima riga è un esempio: sovrascrivila o cancellala.', {'A': 12, 'B': 11, 'C': 11, 'D': 12, 'E': 13, 'F': 13, 'G': 10, 'H': 12, 'I': 34, 'J': 4, 'K': 26, 'L': 14})
intest(ws, 4, 1, ['Data', 'Simbolo', 'Direzione', 'Rischio (€)', 'Risultato (€)', 'Ho seguito il piano?', 'R', 'Esito', 'Nota'])
N = 300
dv1 = DataValidation(type='list', formula1='"long,short"', allow_blank=True); ws.add_data_validation(dv1); dv1.add(f'C5:C{4+N}')
dv2 = DataValidation(type='list', formula1='"sì,no"', allow_blank=True); ws.add_data_validation(dv2); dv2.add(f'F5:F{4+N}')
for r in range(5, 5 + N):
    for col, fmt in (('A', 'DD/MM/YYYY'), ('B', '@'), ('C', '@'), ('D', EUR), ('E', EUR), ('F', '@'), ('I', '@')):
        c = ws[f'{col}{r}']; c.font = f_in; c.number_format = fmt; c.border = bordo
    ws[f'G{r}'] = f'=IF(OR(D{r}="",D{r}=0,E{r}=""),"",E{r}/D{r})'; ws[f'G{r}'].number_format = '+0.00;-0.00;0.00'
    ws[f'H{r}'] = f'=IF(E{r}="","",IF(E{r}>0,"vinto",IF(E{r}<0,"perso","pari")))'
    for col in 'GH': ws[f'{col}{r}'].font = f_lab; ws[f'{col}{r}'].border = bordo
import datetime
es = [datetime.date(2026, 10, 12), 'EURUSD', 'long', 100, 180, 'sì', None, None, 'Esempio: stop 25 pip, 0,40 lotti, target 45 pip']
for i, v in enumerate(es):
    if v is not None: ws.cell(row=5, column=1 + i, value=v)
for r in range(5, 5 + N):
    for col in 'ABCDEFI': ws[f'{col}{r}'].fill = fill_in
R = f'5:{4+N}'
G_, E_, D_, F_ = f'G5:G{4+N}', f'E5:E{4+N}', f'D5:D{4+N}', f'F5:F{4+N}'
ws['K4'] = 'Statistiche'; ws['K4'].font = f_out
stat = [
    ('Trade registrati', f'=COUNT({E_})', '0'),
    ('Vinti', f'=COUNTIF({E_},">0")', '0'),
    ('Win rate', '=IF(L5=0,0,L6/L5)', '0.0%'),
    ('Risultato totale (€)', f'=SUM({E_})', EUR),
    ('Totale in R', f'=SUM({G_})', '+0.00;-0.00;0.00'),
    ('Media R per trade (expectancy)', '=IF(L5=0,0,L9/L5)', '+0.00;-0.00;0.00'),
    ('Guadagno medio (R)', f'=IFERROR(AVERAGEIF({G_},">0"),0)', '0.00'),
    ('Perdita media (R)', f'=IFERROR(-AVERAGEIF({G_},"<0"),0)', '0.00'),
    ('Profit factor', f'=IF(SUMIF({E_},"<0")=0,0,SUMIF({E_},">0")/-SUMIF({E_},"<0"))', '0.00'),
    ('Miglior trade (€)', f'=IF(L5=0,0,MAX({E_}))', EUR),
    ('Peggior trade (€)', f'=IF(L5=0,0,MIN({E_}))', EUR),
    ('Rischio medio (€)', f'=IFERROR(AVERAGE({D_}),0)', EUR),
    ('Piano seguito', f'=IF(COUNTA({F_})=0,0,COUNTIF({F_},"sì")/COUNTA({F_}))', '0%'),
    ('Risultato quando NON segui il piano (€)', f'=SUMIF({F_},"no",{E_})', EUR),
]
for i, (n, f, fmt) in enumerate(stat):
    r = 5 + i
    ws.cell(row=r, column=11, value=n).font = f_lab
    out(ws, f'L{r}', f, fmt)
ws['K20'] = 'Profit factor: guadagni totali ÷ perdite totali.'; ws['K20'].font = f_sot
ws['K21'] = 'Sopra 1 = il sistema guadagna.'; ws['K21'].font = f_sot
ws.freeze_panes = 'A5'
ws['K23'] = AVVISO; ws['K23'].font = f_sot; ws['K23'].alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[23].height = 52

for s in wb.worksheets:
    s.sheet_properties.tabColor = 'E9A10F'
import sys
wb.calculation.fullCalcOnLoad = True
wb.save(sys.argv[1] if len(sys.argv) > 1 else 'Riskly-kit-rischio.xlsx')
